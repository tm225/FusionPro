"""CroisementDocs - croisement de fichiers Excel, CSV, Word, PowerPoint. 100 % hors ligne."""
import os
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog
import pandas as pd

SOURCES = {}   # "fichier :: feuille" -> DataFrame
TEXTES = []    # (lieu, texte) : paragraphes Word, textes de diapos, lignes .txt
MODES = {
    "Éléments communs (A et B)": "inner",
    "Tout A + correspondances de B": "left",
    "Tout A et tout B": "outer",
    "Dans A mais absents de B": "anti_ab",
    "Dans B mais absents de A": "anti_ba",
}


def uniq(cols):
    vus, out = {}, []
    for c in cols:
        c = str(c).strip() or "Colonne"
        n = vus.get(c, 0)
        vus[c] = n + 1
        out.append(c if n == 0 else f"{c}_{n + 1}")
    return out


def table_depuis_lignes(rows):
    if len(rows) < 2:
        return None
    return pd.DataFrame(rows[1:], columns=uniq(rows[0]))


def charger(path):
    nom, ext = os.path.basename(path), os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        for feuille, df in pd.read_excel(path, sheet_name=None, dtype=str).items():
            df.columns = uniq(df.columns)
            SOURCES[f"{nom} :: {feuille}"] = df
    elif ext in (".csv", ".tsv"):
        try:
            df = pd.read_csv(path, sep=None, engine="python", dtype=str, encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(path, sep=None, engine="python", dtype=str, encoding="latin-1")
        df.columns = uniq(df.columns)
        SOURCES[nom] = df
    elif ext == ".docx":
        from docx import Document
        d = Document(path)
        for i, p in enumerate(d.paragraphs, 1):
            if p.text.strip():
                TEXTES.append((f"{nom} :: paragraphe {i}", p.text.strip()))
        for i, t in enumerate(d.tables, 1):
            df = table_depuis_lignes([[c.text.strip() for c in r.cells] for r in t.rows])
            if df is not None:
                SOURCES[f"{nom} :: Tableau {i}"] = df
    elif ext == ".pptx":
        from pptx import Presentation
        for n, slide in enumerate(Presentation(path).slides, 1):
            for sh in slide.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    TEXTES.append((f"{nom} :: diapo {n}", sh.text_frame.text.strip()))
                if getattr(sh, "has_table", False) and sh.has_table:
                    df = table_depuis_lignes([[c.text.strip() for c in r.cells] for r in sh.table.rows])
                    if df is not None:
                        SOURCES[f"{nom} :: diapo {n} tableau"] = df
    elif ext == ".txt":
        with open(path, encoding="utf-8", errors="replace") as f:
            for i, ligne in enumerate(f, 1):
                if ligne.strip():
                    TEXTES.append((f"{nom} :: ligne {i}", ligne.strip()))
    else:
        raise ValueError(f"Format non pris en charge : {ext}")


def cle(serie):
    return serie.fillna("").astype(str).str.strip().str.casefold()


class Dialogue(tk.Toplevel):
    """Choix de source(s), colonne(s) et mode."""

    def __init__(self, parent, titre, deux=True, modes=None):
        super().__init__(parent)
        self.title(titre)
        self.resultat = None
        self.transient(parent)
        self.grab_set()
        self.v = {k: tk.StringVar() for k in ("sa", "ca", "sb", "cb", "mode")}
        self.cb = {}
        noms = list(SOURCES)
        champs = [("sa", "Source A", noms), ("ca", "Colonne clé A", [])]
        if deux:
            champs += [("sb", "Source B", noms), ("cb", "Colonne clé B", [])]
        if modes:
            champs.append(("mode", "Type de croisement", list(modes)))
        for r, (k, lab, vals) in enumerate(champs):
            ttk.Label(self, text=lab).grid(row=r, column=0, sticky="w", padx=8, pady=5)
            c = ttk.Combobox(self, textvariable=self.v[k], values=vals, state="readonly", width=62)
            c.grid(row=r, column=1, padx=8)
            self.cb[k] = c
        self.cb["sa"].bind("<<ComboboxSelected>>", lambda e: self.maj("sa", "ca"))
        if deux:
            self.cb["sb"].bind("<<ComboboxSelected>>", lambda e: self.maj("sb", "cb"))
        if modes:
            self.v["mode"].set(list(modes)[0])
        self.deux, self.modes = deux, modes
        ttk.Button(self, text="Lancer", command=self.ok).grid(row=len(champs), column=1, sticky="e", padx=8, pady=8)
        self.wait_window()

    def maj(self, s, c):
        self.cb[c]["values"] = list(SOURCES[self.v[s].get()].columns)
        self.v[c].set("")

    def ok(self):
        need = ["sa", "ca"] + (["sb", "cb"] if self.deux else [])
        if not all(self.v[k].get() for k in need):
            messagebox.showwarning("Champs manquants", "Veuillez tout renseigner.", parent=self)
            return
        self.resultat = {k: x.get() for k, x in self.v.items()}
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("CroisementDocs — croisement de fichiers (hors ligne)")
        self.geometry("1150x650")
        self.resultat = None
        self.construire_menu()
        corps = ttk.PanedWindow(self, orient="horizontal")
        corps.pack(fill="both", expand=True)
        gauche = ttk.Frame(corps)
        ttk.Label(gauche, text="Sources chargées (double-clic = afficher)").pack(anchor="w", padx=4, pady=2)
        self.liste = tk.Listbox(gauche, selectmode="extended", width=42)
        self.liste.pack(fill="both", expand=True, padx=4)
        self.liste.bind("<Double-1>", self.afficher_source)
        droite = ttk.Frame(corps)
        self.arbre = ttk.Treeview(droite, show="headings")
        sy = ttk.Scrollbar(droite, orient="vertical", command=self.arbre.yview)
        sx = ttk.Scrollbar(droite, orient="horizontal", command=self.arbre.xview)
        self.arbre.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        sy.pack(side="right", fill="y")
        sx.pack(side="bottom", fill="x")
        self.arbre.pack(fill="both", expand=True)
        corps.add(gauche, weight=1)
        corps.add(droite, weight=4)
        self.statut = tk.StringVar(value="Ajoutez des fichiers via le menu Fichier.")
        ttk.Label(self, textvariable=self.statut, relief="sunken", anchor="w").pack(fill="x")

    def construire_menu(self):
        m = tk.Menu(self)
        f = tk.Menu(m, tearoff=0)
        f.add_command(label="Ajouter des fichiers…", command=self.ajouter)
        f.add_command(label="Ajouter un dossier entier…", command=self.ajouter_dossier)
        f.add_command(label="Vider la liste", command=self.vider)
        f.add_separator()
        f.add_command(label="Quitter", command=self.destroy)
        m.add_cascade(label="Fichier", menu=f)
        c = tk.Menu(m, tearoff=0)
        c.add_command(label="Croiser deux sources (jointure / différences)…", command=self.croiser)
        c.add_command(label="Fusionner les sources sélectionnées", command=self.fusionner)
        c.add_command(label="Détecter les doublons…", command=self.doublons)
        m.add_cascade(label="Croisement", menu=c)
        o = tk.Menu(m, tearoff=0)
        o.add_command(label="Rechercher un mot dans tous les fichiers…", command=self.rechercher)
        o.add_command(label="Résumé de la source sélectionnée", command=self.resume)
        m.add_cascade(label="Outils", menu=o)
        r = tk.Menu(m, tearoff=0)
        r.add_command(label="Exporter le résultat (Excel / CSV)…", command=self.exporter)
        m.add_cascade(label="Résultats", menu=r)
        a = tk.Menu(m, tearoff=0)
        a.add_command(label="Mode d'emploi", command=self.aide)
        m.add_cascade(label="Aide", menu=a)
        self.config(menu=m)

    # ---------- fichiers ----------
    def rafraichir(self):
        self.liste.delete(0, "end")
        for n, df in SOURCES.items():
            self.liste.insert("end", f"{n}  [{len(df)} lignes]")
        self.statut.set(f"{len(SOURCES)} tableau(x), {len(TEXTES)} bloc(s) de texte chargés.")

    def charger_chemins(self, chemins):
        for p in chemins:
            try:
                charger(p)
            except Exception as e:
                messagebox.showerror("Erreur", f"{os.path.basename(p)}\n{e}")
        self.rafraichir()

    def ajouter(self):
        self.charger_chemins(filedialog.askopenfilenames(
            title="Choisir des fichiers",
            filetypes=[("Documents", "*.xlsx *.xlsm *.xls *.csv *.tsv *.docx *.pptx *.txt"), ("Tous", "*.*")]))

    def ajouter_dossier(self):
        d = filedialog.askdirectory()
        if d:
            ok = (".xlsx", ".xlsm", ".xls", ".csv", ".tsv", ".docx", ".pptx", ".txt")
            self.charger_chemins([os.path.join(d, f) for f in os.listdir(d)
                                  if f.lower().endswith(ok) and not f.startswith("~$")])

    def vider(self):
        SOURCES.clear()
        TEXTES.clear()
        self.rafraichir()

    def selection(self):
        return [list(SOURCES)[i] for i in self.liste.curselection()]

    # ---------- affichage ----------
    def montrer(self, df, titre):
        self.resultat = df
        self.arbre.delete(*self.arbre.get_children())
        cols = list(df.columns)
        self.arbre["columns"] = cols
        for c in cols:
            self.arbre.heading(c, text=c)
            self.arbre.column(c, width=130, minwidth=60)
        for row in df.head(5000).fillna("").itertuples(index=False):
            self.arbre.insert("", "end", values=[str(x) for x in row])
        extra = " (affichage limité à 5000, l'export contient tout)" if len(df) > 5000 else ""
        self.statut.set(f"{titre} : {len(df)} ligne(s){extra}")

    def afficher_source(self, _=None):
        s = self.selection()
        if s:
            self.montrer(SOURCES[s[0]], s[0])

    def resume(self):
        s = self.selection()
        if not s:
            return messagebox.showinfo("Résumé", "Sélectionnez une source.")
        df = SOURCES[s[0]]
        res = pd.DataFrame({"Colonne": df.columns,
                            "Valeurs remplies": [int(df[c].notna().sum()) for c in df.columns],
                            "Valeurs vides": [int(df[c].isna().sum()) for c in df.columns],
                            "Valeurs distinctes": [int(df[c].nunique()) for c in df.columns]})
        self.montrer(res, f"Résumé de {s[0]}")

    # ---------- croisements ----------
    def croiser(self):
        if len(SOURCES) < 1:
            return messagebox.showinfo("Croisement", "Chargez d'abord des fichiers.")
        d = Dialogue(self, "Croiser deux sources", True, MODES).resultat
        if not d:
            return
        a, b = SOURCES[d["sa"]].copy(), SOURCES[d["sb"]].copy()
        a["_k"], b["_k"] = cle(a[d["ca"]]), cle(b[d["cb"]])
        how = MODES[d["mode"]]
        if how == "anti_ab":
            res = a[~a["_k"].isin(b["_k"])]
        elif how == "anti_ba":
            res = b[~b["_k"].isin(a["_k"])]
        else:
            res = a.merge(b, on="_k", how=how, suffixes=(" (A)", " (B)"))
        self.montrer(res.drop(columns="_k"), d["mode"])

    def fusionner(self):
        s = self.selection()
        if len(s) < 2:
            return messagebox.showinfo("Fusion", "Sélectionnez au moins 2 sources dans la liste.")
        frames = [SOURCES[n].assign(Source=n) for n in s]
        self.montrer(pd.concat(frames, ignore_index=True), "Fusion")

    def doublons(self):
        if not SOURCES:
            return messagebox.showinfo("Doublons", "Chargez d'abord des fichiers.")
        d = Dialogue(self, "Détecter les doublons", False).resultat
        if not d:
            return
        df = SOURCES[d["sa"]].copy()
        df["_k"] = cle(df[d["ca"]])
        res = df[(df["_k"] != "") & df.duplicated("_k", keep=False)].sort_values("_k")
        self.montrer(res.drop(columns="_k"), f"Doublons sur « {d['ca']} »")

    def rechercher(self):
        mot = simpledialog.askstring("Recherche", "Mot ou expression à chercher :", parent=self)
        if not mot:
            return
        kw, out = mot.casefold(), []
        for nom, df in SOURCES.items():
            for c in df.columns:
                masque = df[c].fillna("").astype(str).str.casefold().str.contains(kw, regex=False)
                for i in df.index[masque]:
                    out.append((nom, f"ligne {df.index.get_loc(i) + 2}, colonne {c}", str(df.at[i, c])))
        for lieu, t in TEXTES:
            if kw in t.casefold():
                out.append((lieu.split(" :: ")[0], lieu.split(" :: ")[-1], t))
        self.montrer(pd.DataFrame(out, columns=["Fichier / feuille", "Emplacement", "Texte"]),
                     f"Recherche « {mot} »")

    # ---------- export / aide ----------
    def exporter(self):
        if self.resultat is None or self.resultat.empty:
            return messagebox.showinfo("Export", "Aucun résultat à exporter.")
        p = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                         filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if p:
            try:
                if p.lower().endswith(".csv"):
                    self.resultat.to_csv(p, index=False, encoding="utf-8-sig", sep=";")
                else:
                    self.resultat.to_excel(p, index=False)
                messagebox.showinfo("Export", "Fichier enregistré.")
            except Exception as e:
                messagebox.showerror("Erreur", str(e))

    def aide(self):
        messagebox.showinfo("Mode d'emploi",
            "1. Fichier > Ajouter des fichiers (Excel, CSV, Word, PowerPoint, TXT).\n"
            "   Chaque feuille Excel et chaque tableau Word/PowerPoint devient une source.\n"
            "2. Croisement > choisissez deux sources et la colonne clé de chacune\n"
            "   (ex. nom, matricule, numéro RSU).\n"
            "3. Outils > Recherche dans tous les fichiers, y compris le texte Word/PPT.\n"
            "4. Résultats > Exporter en Excel ou CSV.\n\n"
            "La comparaison ignore les majuscules et les espaces autour des valeurs.\n"
            "Aucune connexion Internet n'est utilisée.")


if __name__ == "__main__":
    App().mainloop()
